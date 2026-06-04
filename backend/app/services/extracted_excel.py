from __future__ import annotations

from datetime import datetime
from pathlib import Path
from uuid import uuid4

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.core.config import settings

HEADER_FILL = PatternFill("solid", fgColor="0F766E")
HEADER_FONT = Font(color="FFFFFF", bold=True)
TITLE_FONT = Font(size=14, bold=True, color="0F172A")


def create_extraction_workbook(data: dict) -> tuple[str, str]:
    export_dir = Path(settings.export_dir)
    export_dir.mkdir(parents=True, exist_ok=True)

    filename = f"image_extraction_{uuid4().hex}.xlsx"
    output_path = export_dir / filename

    workbook = Workbook()
    rows_sheet = workbook.active
    rows_sheet.title = "Extracted Rows"
    _build_rows_sheet(rows_sheet, data.get("rows") or [], data.get("fields") or {})
    _build_fields_sheet(workbook.create_sheet("Fields"), data.get("fields") or {})
    _build_text_sheet(workbook.create_sheet("Raw Text"), data.get("extracted_text") or "")

    workbook.save(output_path)
    return filename, str(output_path)


def _build_rows_sheet(sheet, rows: list[dict[str, str]], fields: dict[str, str]) -> None:
    sheet["A1"] = "Project Excel - Extracted Rows"
    sheet["A1"].font = TITLE_FONT
    sheet["A2"] = f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}"
    sheet["A2"].font = Font(color="475569")

    if not rows and fields:
        rows = [fields]

    columns = _ordered_columns(rows)
    if not columns:
        sheet["A4"] = "No table rows were detected. Check the Fields and Raw Text sheets."
        sheet["A4"].font = Font(italic=True, color="475569")
        _fit_columns(sheet)
        return

    for col_index, column in enumerate(columns, start=1):
        cell = sheet.cell(row=4, column=col_index, value=column)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for row_index, row in enumerate(rows, start=5):
        for col_index, column in enumerate(columns, start=1):
            value = str(row.get(column, "") or "")
            sheet.cell(row=row_index, column=col_index, value=value).alignment = Alignment(vertical="top", wrap_text=True)

    sheet.freeze_panes = "A5"
    sheet.auto_filter.ref = f"A4:{get_column_letter(len(columns))}{max(len(rows) + 4, 4)}"
    _fit_columns(sheet)


def _build_fields_sheet(sheet, fields: dict[str, str]) -> None:
    sheet["A1"] = "Extracted Fields"
    sheet["A1"].font = TITLE_FONT
    sheet["A3"] = "Field"
    sheet["B3"] = "Value"
    for cell in sheet[3]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT

    if fields:
        for row_index, (name, value) in enumerate(fields.items(), start=4):
            sheet.cell(row=row_index, column=1, value=name)
            sheet.cell(row=row_index, column=2, value=str(value or ""))
    else:
        sheet["A4"] = "No standalone fields were detected."
        sheet["A4"].font = Font(italic=True, color="475569")

    sheet.freeze_panes = "A4"
    _fit_columns(sheet)


def _build_text_sheet(sheet, extracted_text: str) -> None:
    sheet["A1"] = "Raw Extracted Text"
    sheet["A1"].font = TITLE_FONT
    sheet["A3"] = extracted_text or "No raw text was returned."
    sheet["A3"].alignment = Alignment(wrap_text=True, vertical="top")
    sheet.column_dimensions["A"].width = 110
    sheet.row_dimensions[3].height = 220


def _ordered_columns(rows: list[dict[str, str]]) -> list[str]:
    columns: list[str] = []
    seen: set[str] = set()
    for row in rows:
        for key in row:
            clean_key = str(key).strip()
            if clean_key and clean_key not in seen:
                seen.add(clean_key)
                columns.append(clean_key)
    return columns


def _fit_columns(sheet) -> None:
    for column_cells in sheet.columns:
        letter = get_column_letter(column_cells[0].column)
        max_length = 0
        for cell in column_cells:
            value = "" if cell.value is None else str(cell.value)
            max_length = max(max_length, min(len(value), 60))
        sheet.column_dimensions[letter].width = max(14, min(max_length + 3, 64))
